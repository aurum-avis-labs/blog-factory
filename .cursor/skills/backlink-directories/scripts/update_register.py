#!/usr/bin/env python3
"""Apply status changes to the backlink register (CSV is the source of truth, XLSX is kept in sync).

Usage
  python3 update_register.py changes.json [--dry-run]
  python3 update_register.py --inline '[["Vemoir","Daily Pings","live","Verified.","https://dailypings.com/p/vemoir","yes"]]'

A change is a list:  [product, platform, status, note, url_or_null, badge_on_site_or_omitted]
  - product / platform are matched case-insensitively; a missing pair is appended as a new row
  - status must be one of STATUSES (see reference/register.md)
  - note is prepended to the Notes column as "Agent <date>: <note> | <older notes>" (capped at 900 chars)
  - url replaces "Listing URL" when given (null keeps the old one)
  - badge is "yes" or "no" and sets "Badge on site" (omit to keep the old value)

The register lives in backlinks/data/register/ (gitignored, never pushed). Override with BACKLINKS_REGISTER_DIR.
A timestamped backup of both files is written to backlinks/data/register/backup/ before every write.
"""
import csv, datetime, json, os, shutil, sys
from pathlib import Path

STATUSES = {
    "live", "submitted", "review", "pending_review", "scheduled", "draft", "planned", "parked",
    "blocked", "blocked_badge", "blocked_captcha", "blocked_other", "blocked_paid",
    "blocked_google_sso", "needs_email_verify", "error",
}
CSV_NAME = "aal-directory-backlinks.csv"
XLSX_NAME = "aal-directory-backlinks.xlsx"


def register_dir() -> Path:
    env = os.environ.get("BACKLINKS_REGISTER_DIR")
    if env:
        return Path(env).expanduser()
    # .cursor/skills/backlink-directories/scripts/update_register.py -> repo root is parents[4]
    return Path(__file__).resolve().parents[4] / "backlinks" / "data" / "register"


def load(csv_path: Path):
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        raise SystemExit(f"{csv_path} has no rows")
    return rows, list(rows[0].keys())


def apply(changes, rows, fields):
    today = datetime.date.today().isoformat()
    touched = []
    for ch in changes:
        if len(ch) < 4:
            raise SystemExit(f"bad change (need at least product, platform, status, note): {ch}")
        product, platform, status, note = ch[:4]
        url = ch[4] if len(ch) > 4 else None
        badge = ch[5] if len(ch) > 5 else None
        if status not in STATUSES:
            raise SystemExit(f"unknown status '{status}' in {ch}; allowed: {sorted(STATUSES)}")
        match = [r for r in rows if r["Product"].lower() == product.lower() and r["Platform"].lower() == platform.lower()]
        if not match:
            new = {k: "" for k in fields}
            new.update({"Product": product, "Platform": platform})
            rows.append(new)
            match = [new]
        for r in match:
            r["Status"] = status
            r["Last update"] = today
            if url:
                r["Listing URL"] = url
            if badge in ("yes", "no"):
                r["Badge on site"] = badge
            old = r.get("Notes", "")
            r["Notes"] = (f"Agent {today}: {note}" + (f" | {old}" if old else ""))[:900]
            touched.append((r["Product"], r["Platform"], status))
    return touched


def write_csv(csv_path: Path, rows, fields):
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def sync_xlsx(xlsx_path: Path, rows):
    """CSV and XLSX keep the same row order, so the sync is positional (duplicate keys are allowed)."""
    if not xlsx_path.exists():
        return False
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    for i, r in enumerate(rows):
        row = i + 2
        values = [r.get(h, "") for h in header]
        if row <= ws.max_row:
            for c, v in enumerate(values, 1):
                ws.cell(row, c).value = v
        else:
            ws.append(values)
    if ws.max_row > len(rows) + 1:
        ws.delete_rows(len(rows) + 2, ws.max_row - (len(rows) + 1))
    wb.save(xlsx_path)
    return True


def main(argv):
    dry = "--dry-run" in argv
    args = [a for a in argv if a != "--dry-run"]
    if not args:
        raise SystemExit(__doc__)
    if args[0] == "--inline":
        changes = json.loads(args[1])
    else:
        with open(args[0], encoding="utf-8") as f:
            changes = json.load(f)
    d = register_dir()
    csv_path, xlsx_path = d / CSV_NAME, d / XLSX_NAME
    if not csv_path.exists():
        raise SystemExit(f"register not found: {csv_path}")
    rows, fields = load(csv_path)
    touched = apply(changes, rows, fields)
    for t in touched:
        print("  ", *t)
    if dry:
        print(f"dry run: {len(touched)} row(s) would change; nothing written")
        return
    backup = d / "backup"
    backup.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(csv_path, backup / f"{CSV_NAME}.{stamp}.bak")
    if xlsx_path.exists():
        shutil.copy2(xlsx_path, backup / f"{XLSX_NAME}.{stamp}.bak")
    write_csv(csv_path, rows, fields)
    synced = sync_xlsx(xlsx_path, rows)
    print(f"APPLIED {len(touched)} row(s); csv written, xlsx {'synced' if synced else 'not found (skipped)'}")


if __name__ == "__main__":
    main(sys.argv[1:])
