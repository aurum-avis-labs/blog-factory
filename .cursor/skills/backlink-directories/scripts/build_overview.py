#!/usr/bin/env python3
"""Build the management overview workbook from the register CSV.

Usage
  python3 build_overview.py [--out FILE] [--notes FILE]

Output (default): backlinks/data/register/overview-YYYY-MM-DD.xlsx  (gitignored)
Sheets: Overview (counts + platform x product matrix), Detail (one row per kept platform x product),
        Dropped platforms (nothing was ever submitted there), Rules & notes.

Rules
  * A platform is KEPT if at least one product has a real listing (live / submitted / review / pending_review /
    scheduled, or a listing that stays hidden until the badge is verified, see HIDDEN_UNTIL_BADGE).
    Kept platforms are shown for ALL products so gaps and their reasons are visible.
  * Every other platform goes to "Dropped platforms" with the reason.

Optional notes file (JSON, default backlinks/data/register/overview-notes.json):
  {"rules": ["free text line", ...],
   "points": {"Platform": "one-line summary shown in the Overview"},
   "blockers": {"Platform|Product": "reason shown in Detail"},
   "owners": {"Platform|Product": "who has to act"},
   "hidden_until_badge": ["Daily Pings", "findly.tools"]}
The workbook uses COUNTIF formulas over the Detail sheet, so run recalc (LibreOffice) if you need cached values.
"""
import collections, csv, datetime, json, os, re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REAL = {"live", "submitted", "review", "pending_review", "scheduled", "scheduled; pending review"}
CATS = ["Live", "Scheduled", "In review", "Waiting for badge", "Planned", "Blocked", "Not attempted"]
COLORS = {"Live": "C6EFCE", "Scheduled": "DDEBF7", "In review": "FFF2CC", "Waiting for badge": "FCE4D6",
          "Planned": "E2E2F0", "Blocked": "F4CCCC", "Not attempted": "EDEDED"}
SHORT = {"Live": "Live", "Scheduled": "Sched.", "In review": "Review", "Waiting for badge": "Badge?",
         "Planned": "Planned", "Blocked": "Blocked", "Not attempted": "-"}


def register_dir() -> Path:
    env = os.environ.get("BACKLINKS_REGISTER_DIR")
    return Path(env).expanduser() if env else Path(__file__).resolve().parents[4] / "backlinks" / "data" / "register"


def first_note(n: str) -> str:
    seg = n.split(" | ")[0]
    seg = re.sub(r"^(Agent|cloud:[\w-]+|delta [\w-]+:?)\s*\d{0,4}-?\d{0,2}-?\d{0,2}:?\s*", "", seg).strip()
    return re.sub(r"^\d{4}-\d\d-\d\d:?\s*", "", seg)[:260]


def main(argv):
    d = register_dir()
    out = Path(argv[argv.index("--out") + 1]) if "--out" in argv else d / f"overview-{datetime.date.today().isoformat()}.xlsx"
    notes_path = Path(argv[argv.index("--notes") + 1]) if "--notes" in argv else d / "overview-notes.json"
    notes = json.load(open(notes_path, encoding="utf-8")) if notes_path.exists() else {}
    hidden = set(notes.get("hidden_until_badge", ["Daily Pings", "findly.tools"]))

    with open(d / "aal-directory-backlinks.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    products = sorted({r["Product"] for r in rows}, key=str.lower)
    reg = collections.defaultdict(dict)
    for r in rows:
        reg[r["Platform"]][r["Product"]] = r

    def has_listing(r):
        return r["Status"] in REAL or (r["Status"] == "blocked_badge" and r["Platform"] in hidden)

    def category(r):
        st = r["Status"]
        if st == "live": return "Live"
        if st.startswith("scheduled"): return "Scheduled"
        if st in ("submitted", "review", "pending_review"): return "In review"
        if has_listing(r): return "Waiting for badge"
        if st in ("planned", "draft"): return "Planned"
        return "Blocked"

    def owner(cat, r):
        key = f"{r['Platform']}|{r['Product']}"
        if key in notes.get("owners", {}): return notes["owners"][key]
        return {"Live": "Nothing to do", "Scheduled": "Nothing to do - goes live on its own",
                "In review": "Waiting for the vendor review", "Waiting for badge": "Dev: add the badge, then verify",
                "Planned": "Agent: submit", "Not attempted": "Decide whether to submit"}.get(cat) or (
            "Owner: decide (paid only, not paying)" if r["Status"] == "blocked_paid" else
            "Owner: solve by hand or skip" if r["Status"] == "blocked_captcha" else "Owner: review")

    keep = sorted([p for p, dd in reg.items() if any(has_listing(r) for r in dd.values())], key=str.lower)
    drop = sorted([p for p in reg if p not in keep], key=str.lower)

    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="305496")

    def header(ws, row, values):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row, c, v)
            cell.font, cell.fill = hdr_font, hdr_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    # ---- Detail
    wd = wb.active
    wd.title = "Detail"
    header(wd, 1, ["Platform", "Product", "Category", "Status", "Listing URL", "Submitted", "Badge needed", "Badge on site", "Note / blocker", "Next step"])
    cat_of = {}
    n = 2
    for p in keep:
        for prod in products:
            r = reg[p].get(prod)
            if r is None:
                cat = "Not attempted"
                vals = [p, prod, cat, "", "", "", "", "", "No register entry for this product on this platform.", owner(cat, {"Platform": p, "Product": prod, "Status": ""})]
            else:
                cat = category(r)
                note = notes.get("blockers", {}).get(f"{p}|{prod}") or first_note(r["Notes"])
                vals = [p, prod, cat, r["Status"], r["Listing URL"], r["Submitted (date)"], r["Badge needed"], r["Badge on site"], note, owner(cat, r)]
            cat_of[(p, prod)] = cat
            for c, v in enumerate(vals, 1):
                cell = wd.cell(n, c, v)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            wd.cell(n, 3).fill = PatternFill("solid", fgColor=COLORS[cat])
            n += 1
    last = n - 1
    for c, w in enumerate([20, 14, 18, 16, 38, 12, 12, 12, 70, 34], 1):
        wd.column_dimensions[get_column_letter(c)].width = w
    wd.freeze_panes = "C2"
    wd.auto_filter.ref = f"A1:J{last}"

    # ---- Overview
    wo = wb.create_sheet("Overview", 0)
    wo["A1"] = "Backlink directory overview"
    wo["A1"].font = Font(bold=True, size=14)
    wo["A2"] = f"Generated {datetime.date.today().isoformat()} from the register. Kept platforms: {len(keep)}. Dropped: {len(drop)}."
    header(wo, 4, ["Category", "Rows"])
    for i, cat in enumerate(CATS, 5):
        wo.cell(i, 1, cat).fill = PatternFill("solid", fgColor=COLORS[cat])
        wo.cell(i, 2, f'=COUNTIF(Detail!$C$2:$C${last},A{i})')
    wo.cell(5 + len(CATS), 1, "Total").font = Font(bold=True)
    wo.cell(5 + len(CATS), 2, f"=SUM(B5:B{4 + len(CATS)})").font = Font(bold=True)

    m0 = 7 + len(CATS)
    header(wo, m0, ["Platform"] + products + ["Summary"])
    for i, p in enumerate(keep, m0 + 1):
        wo.cell(i, 1, p).font = Font(bold=True)
        for j, prod in enumerate(products, 2):
            cat = cat_of[(p, prod)]
            cell = wo.cell(i, j, SHORT[cat])
            cell.fill = PatternFill("solid", fgColor=COLORS[cat])
            cell.alignment = Alignment(horizontal="center")
        wo.cell(i, len(products) + 2, notes.get("points", {}).get(p, "")).alignment = Alignment(wrap_text=True, vertical="top")
    wo.column_dimensions["A"].width = 24
    for j in range(2, len(products) + 2):
        wo.column_dimensions[get_column_letter(j)].width = 12
    wo.column_dimensions[get_column_letter(len(products) + 2)].width = 70

    # ---- Dropped
    wx = wb.create_sheet("Dropped platforms")
    header(wx, 1, ["Platform", "Products in register", "Why nothing was submitted"])
    for i, p in enumerate(drop, 2):
        rs = list(reg[p].values())
        c = collections.Counter(r["Status"] for r in rs)
        reason = notes.get("points", {}).get(p) or first_note(max(rs, key=lambda r: len(r["Notes"]))["Notes"])
        wx.cell(i, 1, p)
        wx.cell(i, 2, ", ".join(f"{k} x{v}" for k, v in c.items()))
        wx.cell(i, 3, reason).alignment = Alignment(wrap_text=True, vertical="top")
    for c, w in enumerate([24, 30, 100], 1):
        wx.column_dimensions[get_column_letter(c)].width = w

    # ---- Rules
    wr = wb.create_sheet("Rules & notes")
    lines = ["Kept = the platform holds at least one real submission: live, scheduled, in review, or accepted but hidden until our badge is verified.",
             "Dropped = nothing was ever submitted there (blocked, parked, planned or an unsubmitted draft); reasons are in 'Dropped platforms'.",
             "No payment is ever made; paid tiers, skip-the-line and Premium offers are declined.",
             "Not attempted = the product was not submitted on a platform that was only used for some products."] + notes.get("rules", [])
    for i, t in enumerate(lines, 1):
        wr.cell(i, 1, "- " + t).alignment = Alignment(wrap_text=True, vertical="top")
    wr.column_dimensions["A"].width = 140

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}  (kept {len(keep)}, dropped {len(drop)}, detail rows {last - 1})")


if __name__ == "__main__":
    main(sys.argv[1:])
